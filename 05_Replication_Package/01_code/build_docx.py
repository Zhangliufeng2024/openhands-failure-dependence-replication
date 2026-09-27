# -*- coding: utf-8 -*-
"""Build the submittable .docx from the manuscript Markdown source.

Renders inline bold/italic/code markup, sub- and super-scripts (_{...}, ^{...}),
display equations (indented blocks), real Word tables with booktabs-style rules
and merged multi-row headers, numbered lists, figures with captions, continuous
line numbers, page numbers, and hyperlinked URLs.

Usage:  python build_docx.py [source.md] [output.docx]
Defaults are resolved relative to this script's location.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..'))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '01_Manuscript', 'Manuscript_Paper3_v3.7_source.md')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, '01_Manuscript', 'Manuscript_Paper3_v3.7.docx')
FIGDIR = os.path.join(ROOT, '02_Figures')

SERIF, MONO, EASIA = 'Times New Roman', 'Consolas', '宋体'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)   # A4
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(1)

st = doc.styles['Normal']
st.font.name = SERIF
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), EASIA)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# continuous line numbering (for review)
ln = OxmlElement('w:lnNumType')
ln.set(qn('w:countBy'), '1')
ln.set(qn('w:restart'), 'continuous')
ln.set(qn('w:distance'), '360')
sec._sectPr.append(ln)

def set_font(run, name=SERIF, size=None, bold=None, italic=None):
    run.font.name = name
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'), EASIA)

# ---- footer page number ---------------------------------------------------
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for el, attrs, txt in [('w:fldChar', {'w:fldCharType': 'begin'}, None),
                       ('w:instrText', {'xml:space': 'preserve'}, 'PAGE'),
                       ('w:fldChar', {'w:fldCharType': 'end'}, None)]:
    e = OxmlElement(el)
    for k, v in attrs.items(): e.set(qn(k), v)
    if txt: e.text = txt
    r._r.append(e)
set_font(r, size=9)

# ---- inline markup --------------------------------------------------------
TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*|_\{[^}]*\}|\^\{[^}]*\}|https?://\S+)')

def add_sub_sup(p, text, size=None, bold=False, italic=False, sub=True):
    # subscript/superscript content may itself carry *italic* markup
    for part in re.split(r'(\*[^*]+?\*)', text):
        if not part:
            continue
        if part.startswith('*') and part.endswith('*') and len(part) >= 2:
            rr = p.add_run(part[1:-1]); set_font(rr, size=size, bold=bold, italic=True)
        else:
            rr = p.add_run(part); set_font(rr, size=size, bold=bold, italic=italic)
        rr.font.subscript = sub
        rr.font.superscript = not sub

def add_rich(p, text, size=None, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            set_font(p.add_run(part[2:-2]), size=size, bold=True, italic=italic)
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            set_font(p.add_run(part[1:-1]), name=MONO, size=(size or 11) - 1.5, bold=bold, italic=italic)
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**') and len(part) >= 2:
            set_font(p.add_run(part[1:-1]), size=size, italic=True, bold=bold)
        elif part.startswith('_{') and part.endswith('}'):
            add_sub_sup(p, part[2:-1], size=size, bold=bold, italic=italic, sub=True)
        elif part.startswith('^{') and part.endswith('}'):
            add_sub_sup(p, part[2:-1], size=size, bold=bold, italic=italic, sub=False)
        elif part.startswith('http'):
            rr = p.add_run(part); set_font(rr, size=size, bold=bold, italic=italic)
            rid = doc.part.relate_to(part, RT.HYPERLINK, is_external=True)
            h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
            rr._r.getparent().remove(rr._r)
            c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1')
            rr._element.get_or_add_rPr().append(c)
            h.append(rr._r); p._p.append(h)
        else:
            set_font(p.add_run(part), size=size, bold=bold, italic=italic)

# ---- structural helpers ----------------------------------------------------
def heading(text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=13 if level == 1 else 11.5, bold=True)
    if level == 1:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)

def body(text):
    p = doc.add_paragraph()
    add_rich(p, text)

def caption(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_rich(p, text, size=9.5, italic=True)

def figure_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    add_rich(p, text, size=9.5, italic=True)

def table_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    add_rich(p, text, size=8.5)

def display_equation(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    add_rich(p, text, size=11.5)

def add_figure(fn):
    doc.add_picture(os.path.join(FIGDIR, fn), width=Inches(6.27))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

def set_cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    el = borders.find(qn('w:' + edge))
    if el is None:
        el = OxmlElement('w:' + edge); borders.append(el)
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')

def split_row(line):
    """Split a space-aligned row into (start_col, text) cells."""
    cells, pos = [], 0
    for m in re.finditer(r'\s{2,}', line):
        cells.append((pos, line[pos:m.start()].strip()))
        pos = m.end()
    cells.append((pos, line[pos:].strip()))
    return cells

def col_of(start, bounds):
    return max(i for i, b in enumerate(bounds) if start + 1 >= b)

def build_table(groups):
    hdr_lines = groups[0]
    # column boundaries from the last header line; column 0 always exists
    bounds = sorted(set([0] + [c[0] for c in split_row(hdr_lines[-1]) if c[1] != '' and c[0] > 0]))
    ncols = len(bounds)

    def assign(line):
        row = [''] * ncols
        for start, txt in split_row(line):
            if txt:
                i = col_of(start, bounds)
                row[i] = (row[i] + '\n' + txt).strip() if row[i] else txt
        return row

    # header: one row per header line, merging spans for panel titles
    hdr_rows = [assign(h) for h in hdr_lines]
    # data rows + notes (a non-indented single-cell line is a table note)
    data, notes = [], []
    for g in groups[1:]:
        for ln in g:
            ncells = sum(1 for _, t in split_row(ln) if t)
            if not ln.startswith(' ') and ncells < 2:
                if ln.strip():
                    notes.append(ln.strip())
                continue
            row = assign(ln)
            if ln.startswith(' ') and data:
                prev = data[-1]
                for i, v in enumerate(row):
                    if v:
                        prev[i] = (prev[i] + '\n' + v).strip() if prev[i] else v
            else:
                data.append(row)

    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for hr in hdr_rows:
        row = tbl.add_row()
        for i, v in enumerate(hr):
            if i >= ncols: break
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, v, size=9, bold=True)
    # merge spans for upper header rows (e.g. Table 5's tool-name row)
    for ri, hr in enumerate(hdr_rows[:-1]):
        starts = [(c[0], c[1]) for c in split_row(hdr_lines[ri]) if c[1]]
        for si, (st, txt) in enumerate(starts):
            i0 = col_of(st, bounds)
            if si + 1 < len(starts):
                i1 = col_of(starts[si + 1][0], bounds) - 1
            else:
                i1 = ncols - 1
            if i1 > i0:
                tbl.cell(ri, i0).merge(tbl.cell(ri, i1))
    for row in data:
        trow = tbl.add_row()
        for i, v in enumerate(row):
            if i >= ncols: break
            p = trow.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, v, size=9)
    # booktabs rules
    for cell in tbl.rows[0].cells:
        set_cell_border(cell, 'top', 12)
    for cell in tbl.rows[len(hdr_rows) - 1].cells:
        set_cell_border(cell, 'bottom', 6)
    for cell in tbl.rows[-1].cells:
        set_cell_border(cell, 'bottom', 12)
    for nt in notes:
        table_note(nt)

# ---- main parse loop -------------------------------------------------------
raw = open(SRC, encoding='utf-8').read()
lines = raw.split('\n')
i, n = 0, len(lines)

def is_sep(ln):
    s = ln.strip()
    return bool(s) and set(s) <= set('-=')

def para_until_blank(i):
    out = []
    while i < n and lines[i].strip() and not is_sep(lines[i]) and not lines[i].startswith('    '):
        out.append(lines[i].strip())
        i += 1
    return ' '.join(out), i

# title: leading non-blank lines until 'Abstract'
title_lines = []
while i < n and lines[i].strip() != 'Abstract':
    if lines[i].strip() and not is_sep(lines[i]):
        title_lines.append(lines[i].strip())
    i += 1

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run('\n'.join(title_lines)); set_font(r, size=16, bold=True)
for txt, sz in [('Liufeng Zhang^{a,*}, Ke Zhuang^{a}', 11),
                ('^{a} College of Science and Technology, Ningbo University, Ningbo 315300, Zhejiang Province, China', 10),
                ('*Corresponding author: zhangliufeng@nbu.edu.cn', 10)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, txt, size=sz)

stats = {'headings': 0, 'tables': 0, 'figures': 0, 'paras': 0}

while i < n:
    line = lines[i]
    s = line.strip()

    if not s or is_sep(line):
        i += 1; continue

    if s.startswith('[FIGURE:'):
        add_figure(s[len('[FIGURE:'):-1].strip())
        stats['figures'] += 1
        i += 1
        if i < n and re.match(r'^Figure \d+\.', lines[i].strip()):
            cap, i = para_until_blank(i)
            figure_note(cap)
        continue

    # a real table caption is followed by a separator line
    if re.match(r'^Table \d+\.', s) and i + 1 < n and is_sep(lines[i + 1]):
        cap, i = para_until_blank(i)
        caption(cap)
        groups, cur = [], []
        while i < n and lines[i].strip():
            ln = lines[i]
            if is_sep(ln):
                if cur: groups.append(cur); cur = []
                i += 1; continue
            if re.match(r'^(Table|Figure) \d+\.', ln.strip()) and not ln.startswith(' '):
                break
            cur.append(ln.rstrip()); i += 1
        if cur: groups.append(cur)
        if groups:
            build_table(groups)
            stats['tables'] += 1
        continue

    if s in ('Abstract', 'References') or s.startswith('Appendix '):
        heading(s, 1); stats['headings'] += 1; i += 1; continue

    if re.match(r'^\d+\.\d+\s+[A-Z]', s) and len(s) < 90 and not s.rstrip().endswith(('.', ',', ';')):
        heading(s, 2); stats['headings'] += 1; i += 1; continue

    if re.match(r'^\d+\s+[A-Z]', s) and len(s) < 90 and not s.rstrip().endswith(('.', ',', ';')):
        heading(s, 1); stats['headings'] += 1; i += 1; continue

    if line.startswith('    '):
        eq = []
        while i < n and lines[i].startswith('    '):
            eq.append(lines[i].strip()); i += 1
        display_equation(' '.join(eq))
        continue

    if re.match(r'^\d+\.\s', s):
        item, i = para_until_blank(i)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.32)
        add_rich(p, item)
        stats['paras'] += 1
        continue

    txt, i = para_until_blank(i)
    if txt.startswith('Keywords:'):
        p = doc.add_paragraph()
        add_rich(p, txt)
        if p.runs: p.runs[0].bold = True
    else:
        body(txt)
    stats['paras'] += 1

doc.core_properties.title = 'Two Scales of Failure Dependence in LLM Agent Trajectories'
doc.core_properties.subject = 'Empirical Software Engineering special issue submission'
doc.core_properties.author = 'Liufeng Zhang, Ke Zhuang'
doc.core_properties.comments = 'Built from ' + os.path.basename(SRC)

doc.save(OUT)
print('saved:', OUT)
print('stats:', stats, '| tables:', len(doc.tables), 'images:', len(doc.inline_shapes))
